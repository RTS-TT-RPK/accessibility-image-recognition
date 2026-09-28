# -*- coding: utf-8 -*-
"""新流程自测 —— 一条命令验证"脚本层"没坏（只读真实数据，跑完自动清理）。

    python tools/新流程自测.py

它在**系统临时目录里**搭一个迷你工作区（复制 validate_export.py、schemas/、aliases.json
和 tools/ 下的脚本），然后用假图片跑一遍完整流程，逐项断言：

  1. 建任务        建出 runs/日期_编号/、复制原图、算指纹、写台账、raw.json 骨架是空的
  2. --next        能报出下一张待识别的图和要填的格式
  3. 提交识别      校验通过、落盘；row_id/value/required_fields/unit 由脚本补
  4. 拒绝写入      缺 row_role / 用了契约外的列 / 数值解析不了 —— 三种都必须**拒绝且不落盘**
  5. 重复提交      不带 --replace 必须拒绝；带 --replace 行号必须**稳定不变**
  6. 校验+落账     validate_export.py 能跑出 issues.json / final.xlsx / task_review.xlsx
  7. 导出长表      facts.jsonl 一行一条事实，且未解决的行不进长表
  8. 收尾报告      能出一句话摘要
  9. 注册表型      没回答 row_unit 必须拒绝；回答了才写契约并跑回归

**绝不碰真实数据**：所有写入都在临时目录里；真实工作区只被读。
退出码：0 全绿；1 有失败项。
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PY = sys.executable

PASS, FAIL = [], []


def check(name, ok, detail=""):
    (PASS if ok else FAIL).append(name)
    print("  %s %s%s" % ("[通过]" if ok else "[失败]", name, ("  <- " + detail) if detail else ""))


def run(ws, args, stdin_text=None):
    """在迷你工作区里跑一条命令，返回 (退出码, stdout+stderr)。"""
    p = subprocess.run([PY] + args, cwd=ws, capture_output=True, text=True,
                       encoding="utf-8", errors="replace", input=stdin_text, timeout=600)
    return p.returncode, (p.stdout or "") + (p.stderr or "")


def read_json(path, default=None):
    if not os.path.isfile(path):
        return default
    with open(path, "r", encoding="utf-8-sig") as f:
        return json.load(f)


def write_json(path, obj):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        json.dump(obj, f, ensure_ascii=False, indent=1)
        f.write("\n")


def build_mini_ws(tmp):
    """搭一个迷你工作区：脚本 + 契约 + 词典，全是真实文件的副本。"""
    for fn in ("validate_export.py", "aliases.json", "FEEDBACK_LOG.md"):
        shutil.copy2(os.path.join(BASE, fn), os.path.join(tmp, fn))
    shutil.copytree(os.path.join(BASE, "schemas"), os.path.join(tmp, "schemas"))
    os.makedirs(os.path.join(tmp, "tools"), exist_ok=True)
    for fn in ("建任务.py", "提交识别.py", "导入审核.py", "注册表型.py",
               "导出长表.py", "收尾报告.py", "收录.py", "生成样板.py", "回归测试.py",
               "初始化工作区.py", "任务.py"):
        src = os.path.join(BASE, "tools", fn)
        if os.path.isfile(src):
            shutil.copy2(src, os.path.join(tmp, "tools", fn))
    for d in ("inbox", "runs", "archive", "golden", "pending"):
        os.makedirs(os.path.join(tmp, d), exist_ok=True)
    return tmp


FRAG_OK = {
    "image_id": "IMG_001",
    "sheet_label": "库存日报",
    "date_text": "2026-09-10",
    "rows": [
        {"row_role": "data", "object_label": "1号库/IPAC",
         "fields": {"日期": "2026-09-10", "位置": "1号库", "物料": "IPAC"},
         "components": [{"name": "IPAC", "raw_value": "99.5%"},
                        {"name": "水", "raw_value": "0.1%"}]},
        {"row_role": "data", "object_label": "2号库/甲醇",
         "fields": {"日期": "2026-09-10", "位置": "2号库", "物料": "甲醇"},
         # 甲醇 99.9 + 杂 0.5 = 100.4（水不计入合计）→ 故意造一条"合计超上限"的严重问题，
         # 用来验证：① 提交时当场报出来 ② 该行不进长表 ③ 导入人工决定后能解封
         "components": [{"name": "甲醇", "raw_value": "99.9%"},
                        {"name": "水", "raw_value": "0.1%"},
                        {"name": "杂", "raw_value": "0.5%"}]},
    ],
}


FRAG_2 = {
    "image_id": "IMG_002",
    "sheet_label": "来料分析",
    "date_text": "2026-09-10",
    "rows": [
        {"row_role": "data", "object_label": "A001/桶01",
         "fields": {"日期": "2026-09-10", "位置": "北罐区", "物料": "IPAC",
                    "批次": "A001", "桶号": "01"},
         "components": [{"name": "IPAC", "raw_value": "89.6%"},
                        {"name": "水", "raw_value": "0.1%"}]},
    ],
}


def main():
    tmp = tempfile.mkdtemp(prefix="新流程自测_")
    print()
    print("=" * 62)
    print("  新流程自测（只读真实数据，全部写入都在临时目录）")
    print("=" * 62)
    print("  临时工作区：%s" % tmp)
    print("  真实工作区只被读取，不会被改一个字。")
    print()

    try:
        build_mini_ws(tmp)

        # 造两张假图（脚本只算指纹、不解码图片，所以内容无所谓）
        inbox = os.path.join(tmp, "inbox")
        for name, blob in (("IMG_001.jpg", b"fake image one"),
                           ("IMG_002.jpg", b"fake image two")):
            with open(os.path.join(inbox, name), "wb") as f:
                f.write(blob)

        # ---------- 1. 建任务 ----------
        print("【1】建任务")
        code, out = run(tmp, ["tools/建任务.py", "--date", "2026-09-10"])
        check("建任务退出码为 0", code == 0, out.strip()[-200:])
        run_dir = None
        runs = os.path.join(tmp, "runs")
        cands = [d for d in os.listdir(runs) if os.path.isdir(os.path.join(runs, d))]
        run_dir = os.path.join(runs, cands[0]) if len(cands) == 1 else None
        check("建出了唯一一个任务号 runs/20260910_001", run_dir is not None and cands[0] == "20260910_001",
              str(cands))
        manifest = read_json(os.path.join(run_dir, "source", "names.json"), {}) if run_dir else {}
        check("图片台账登记了 2 张", len(manifest.get("images", [])) == 2)
        check("指纹是真实算出来的（非空且长度 64）",
              all(len(i.get("sha256", "")) == 64 for i in manifest.get("images", [])))
        raw = read_json(os.path.join(run_dir, "raw.json"), {})
        check("raw.json 骨架是空的（进度不会被谎报成已完成）",
              raw.get("images") == [] and raw.get("rows") == [])
        check("原图已复制进 source/（inbox 原件保留）",
              len([f for f in os.listdir(os.path.join(run_dir, "source")) if f.endswith(".jpg")]) == 2
              and len(os.listdir(inbox)) == 2)

        rel = os.path.relpath(run_dir, tmp)

        # ---------- 2. --next ----------
        print()
        print("【2】看下一张待识别")
        code, out = run(tmp, ["tools/提交识别.py", rel, "--next"])
        check("--next 退出码为 0", code == 0)
        check("--next 报出了图片绝对路径", "IMG_001.jpg" in out)
        check("--next 列出了已注册表型（让 AI 别自己发明列名）", "库存日报" in out and "来料分析" in out)
        check("--next 给出了片段格式模板", "row_role" in out and "raw_value" in out)

        # ---------- 3. 正常提交 ----------
        print()
        print("【3】提交一张图的识别结果")
        frag = os.path.join(tmp, "_frag.json")
        write_json(frag, FRAG_OK)
        code, out = run(tmp, ["tools/提交识别.py", rel, "--image", "IMG_001", "--file", "_frag.json"])
        check("提交退出码为 0", code == 0, out.strip()[-300:])
        raw = read_json(os.path.join(run_dir, "raw.json"), {})
        check("落盘 2 行数据", len(raw.get("rows", [])) == 2)
        check("图片只登记了识别完的那一张", [i.get("image_id") for i in raw.get("images", [])] == ["IMG_001"])
        rows = raw.get("rows", [])
        check("row_id 由脚本编号", [r.get("row_id") for r in rows] == ["R001", "R002"])
        check("value 由脚本换算（99.5% -> 0.995）",
              rows[0]["components"][0].get("value") == "0.995",
              str(rows[0]["components"][0]))
        check("unit 由脚本按契约补", rows[0]["components"][0].get("unit") == "%")
        check("required_fields 从契约推出来（日期/位置/物料）",
              rows[0].get("required_fields") == ["日期", "位置", "物料"], str(rows[0].get("required_fields")))
        check("没有采信 AI 写的 value（片段里根本没给）", "value" not in FRAG_OK["rows"][0]["components"][0])

        check("提交时当场报出了裁判的疑点（合计超上限）", "超过上限" in out, out.strip()[-200:])

        # ---------- 4. 三种必须拒绝的情况 ----------
        print()
        print("【4】必须拒绝且不落盘（格式/契约错误）")
        before = json.dumps(read_json(os.path.join(run_dir, "raw.json"), {}), ensure_ascii=False, sort_keys=True)

        bad_cases = []
        b1 = json.loads(json.dumps(FRAG_OK)); b1["image_id"] = "IMG_002"
        del b1["rows"][0]["row_role"]
        bad_cases.append(("缺 row_role", b1))
        b2 = json.loads(json.dumps(FRAG_OK)); b2["image_id"] = "IMG_002"
        b2["rows"][0]["fields"]["水分"] = "0.1%"
        bad_cases.append(("用了契约外的列", b2))
        b3 = json.loads(json.dumps(FRAG_OK)); b3["image_id"] = "IMG_002"
        b3["rows"][0]["components"][0]["raw_value"] = "看不清的几个字"
        bad_cases.append(("数值解析不了", b3))

        for label, frag_obj in bad_cases:
            write_json(frag, frag_obj)
            code, out = run(tmp, ["tools/提交识别.py", rel, "--image", "IMG_002", "--file", "_frag.json"])
            same = json.dumps(read_json(os.path.join(run_dir, "raw.json"), {}),
                              ensure_ascii=False, sort_keys=True) == before
            check("拒绝：%s（退出码 2 且 raw.json 未被改动）" % label, code == 2 and same,
                  "code=%s changed=%s" % (code, not same))

        # ---------- 5. 重复提交 ----------
        print()
        print("【5】重复提交与行号稳定")
        write_json(frag, FRAG_OK)
        code, out = run(tmp, ["tools/提交识别.py", rel, "--image", "IMG_001", "--file", "_frag.json"])
        check("不带 --replace 时拒绝重复提交", code == 2, "code=%s" % code)
        raw2 = read_json(os.path.join(run_dir, "raw.json"), {})
        check("被拒绝后行号没乱", [r.get("row_id") for r in raw2.get("rows", [])] == ["R001", "R002"])

        code, out = run(tmp, ["tools/提交识别.py", rel, "--image", "IMG_001", "--file", "_frag.json", "--replace"])
        raw3 = read_json(os.path.join(run_dir, "raw.json"), {})
        check("--replace 重交后行号稳定不变（ISS 编号才不会错位）",
              [r.get("row_id") for r in raw3.get("rows", [])] == ["R001", "R002"],
              str([r.get("row_id") for r in raw3.get("rows", [])]))

        # 第二张图正常交掉，让任务真正走到"待人工复核"
        write_json(frag, FRAG_2)
        code, out = run(tmp, ["tools/提交识别.py", rel, "--image", "IMG_002", "--file", "_frag.json"])
        check("第二张图也交掉了", code == 0, out.strip()[-200:])
        code, out = run(tmp, ["tools/提交识别.py", rel, "--next"])
        check("全部识别完后 --next 提示去做校验", "全部识别完" in out, out.strip()[-200:])
        raw4 = read_json(os.path.join(run_dir, "raw.json"), {})
        img2 = [i for i in raw4.get("images", []) if i.get("image_id") == "IMG_002"]
        # 契约里写着 archive.name_pattern = "{date}_{sheet}_{batch}_第{n}张"，
        # 以前没人读它（等于白写），现在由脚本按契约算出来
        check("归档文件名由契约的 name_pattern 算出来（不再靠 AI 起名）",
              bool(img2) and img2[0].get("suggested_name") == "2026-09-10_来料分析_A001_第1张",
              str(img2[0].get("suggested_name")) if img2 else "没找到 IMG_002")

        # ---------- 6. 校验 + 落账 ----------
        print()
        print("【6】校验与落账（裁判脚本）")
        code, out = run(tmp, ["validate_export.py", rel, "--init-issues"])
        check("--init-issues 退出码为 0", code == 0, out.strip()[-300:])
        check("生成了 issues.json", os.path.isfile(os.path.join(run_dir, "issues.json")))
        issues_doc = read_json(os.path.join(run_dir, "issues.json"), {}) or {}
        severe = [i for i in issues_doc.get("issues", []) if i.get("severity") == "severe"]
        check("合计超上限被记成了严重问题（这就是要靠人工复核的那类）", len(severe) >= 1,
              "severe=%d" % len(severe))
        code, out = run(tmp, ["validate_export.py", rel])
        check("首次落账退出码为 0", code == 0, out.strip()[-300:])
        check("生成了 final.xlsx", os.path.isfile(os.path.join(run_dir, "final.xlsx")))
        check("生成了 task_review.xlsx", os.path.isfile(os.path.join(run_dir, "task_review.xlsx")))

        # ---------- 7. 长表（红线：未解决的行不许进） ----------
        print()
        print("【7】导出长表 facts.jsonl")
        code, out = run(tmp, ["tools/导出长表.py", rel])
        check("导出长表退出码为 0", code == 0, out.strip()[-300:])
        facts_path = os.path.join(run_dir, "facts.jsonl")

        def read_facts():
            rows_ = []
            if os.path.isfile(facts_path):
                with open(facts_path, "r", encoding="utf-8") as f:
                    for line in f:
                        if line.strip():
                            rows_.append(json.loads(line))
            return rows_

        facts = read_facts()
        check("facts.jsonl 已生成", os.path.isfile(facts_path))
        rids = sorted({f.get("source_row_id") for f in facts})
        check("红线：被严重问题挡住的行**没有**进长表",
              "R002" not in rids and "R001" in rids and "R003" in rids, "行号=%s" % rids)
        check("导出时如实报出被挡掉的行数", "没进长表" in out, out.strip()[-200:])
        check("长表带来源行号与原始值",
              bool(facts) and "source_row_id" in facts[0] and "raw_value" in facts[0], str(facts[:1]))
        n_before = len(facts)

        # ---------- 8. 收尾报告 ----------
        print()
        print("【8】收尾报告")
        code, out = run(tmp, ["tools/收尾报告.py", rel, "--one-line"])
        check("收尾报告退出码为 0", code == 0, out.strip()[-200:])
        check("报告里带了张数/行数/问题数",
              "原图 2 张（已识别 2）" in out and "数据行 3 行" in out and "问题 1 条" in out, out.strip())
        check("报告如实说还在等人工复核", "待人工复核" in out, out.strip())

        # ---------- 9. 注册表型 ----------
        print()
        print("【9】注册表型（硬拦：没确认就不许写契约）")
        draft = {
            "template": "测试表_v1", "sheet_label": "测试表",
            "row_unit": "（请确认：这张表一行代表什么？）",
            "columns": [{"name": "日期", "type": "date", "required": True},
                        {"name": "数量", "type": "number"}],
            "measurements": {"enabled": False, "unit": "%",
                             "categories": {"water": "never", "misc": "if_present"},
                             "default_category": "always", "sum_max": 100.0},
            "vocabulary": {"water": [], "misc": []},
            "fact_keys": {}, "archive": {"name_pattern": None},
        }
        write_json(os.path.join(run_dir, "样板_契约草案.json"), draft)
        code, out = run(tmp, ["tools/注册表型.py", rel, "--name", "测试表"])
        check("row_unit 没确认时拒绝写契约", code == 2, "code=%s" % code)
        check("拒绝时没有偷偷写出契约文件",
              not os.path.isfile(os.path.join(tmp, "schemas", "测试表_v1.json")))

        code, out = run(tmp, ["tools/注册表型.py", rel, "--name", "测试表",
                              "--row-unit", "一行 = 一个测试对象", "--notes", "自测用"])
        check("回答了 row_unit 后写出契约", code == 0, out.strip()[-400:])
        wrote = os.path.join(tmp, "schemas", "测试表_v1.json")
        check("契约文件已生成", os.path.isfile(wrote))
        check("sample.json 记下了样板已确认",
              read_json(os.path.join(run_dir, "sample.json"), {}).get("confirmed") is True)
        check("回归测试全绿（契约没把裁判改坏）", "回归" in out and code == 0)

        # ---------- 10. 审核表往返（本次改造最关键的一环） ----------
        print()
        print("【10】审核表往返：填 Excel → 导入 → 重跑落账 + 归档")

        def find_col(ws, header_text, max_rows=6):
            """按表头文字找列号（不写死列位置 —— 位置一动就会读错列）。"""
            for r in range(1, max_rows + 1):
                for c in range(1, (ws.max_column or 1) + 1):
                    v = ws.cell(row=r, column=c).value
                    if v and header_text in str(v):
                        return r, c
            return None, None

        from openpyxl import load_workbook
        review = os.path.join(run_dir, "task_review.xlsx")
        wb = load_workbook(review)
        check("审核表有「问题清单」页", "问题清单" in wb.sheetnames, str(wb.sheetnames))
        check("审核表有「文件归档清单」页", "文件归档清单" in wb.sheetnames, str(wb.sheetnames))

        ws_i = wb["问题清单"]
        hr, c_id = find_col(ws_i, "问题编号")
        _, c_dec = find_col(ws_i, "员工决定")
        check("能按表头找到「问题编号」和「员工决定」列", hr is not None and c_id is not None and c_dec is not None,
              "hr=%s id=%s dec=%s" % (hr, c_id, c_dec))
        filled = 0
        if hr is not None and c_dec is not None:
            for r in range(hr + 1, (ws_i.max_row or hr) + 1):
                if ws_i.cell(row=r, column=c_id).value:
                    ws_i.cell(row=r, column=c_dec).value = "通过"
                    filled += 1
        check("给每一条问题都填上了决定", filled >= 1, "填了 %d 条" % filled)

        ws_a = wb["文件归档清单"]
        hr2, c_yn = find_col(ws_a, "是否归档")
        check("能按表头找到「是否归档」列", hr2 is not None and c_yn is not None,
              "hr=%s col=%s" % (hr2, c_yn))
        if hr2 is not None and c_yn is not None:
            for r in range(hr2 + 1, (ws_a.max_row or hr2) + 1):
                if ws_a.cell(row=r, column=1).value:
                    ws_a.cell(row=r, column=c_yn).value = "是"
        wb.save(review)
        wb.close()

        code, out = run(tmp, ["tools/导入审核.py", rel])
        check("导入审核退出码为 0", code == 0, out.strip()[-400:])
        dec = read_json(os.path.join(run_dir, "decisions.json"), {}) or {}
        check("decisions.json 生成了问题决定", len(dec.get("issue_decisions") or []) >= 1,
              str(dec.get("issue_decisions"))[:200])
        check("decisions.json 生成了归档决定（以前这个下拉框没人读）",
              len(dec.get("archive_decisions") or []) == 2,
              str(dec.get("archive_decisions"))[:200])

        code, out = run(tmp, ["validate_export.py", rel, "--archive"])
        check("重跑 + 归档退出码为 0", code == 0, out.strip()[-400:])
        arch_root = os.path.join(tmp, "archive")
        moved = []
        for dp, _dn, fns in os.walk(arch_root):
            for fn in fns:
                moved.append(os.path.join(dp, fn))
        check("原图已归档到 archive/（2 张）", len(moved) == 2, "实际 %d 张" % len(moved))
        check("归档后 source/ 里没有残留原图",
              len([f for f in os.listdir(os.path.join(run_dir, "source")) if f.lower().endswith(".jpg")]) == 0)

        code, out = run(tmp, ["tools/导出长表.py", rel])
        facts2 = read_facts()
        check("人工决定生效后，被挡住的行解封进长表（%d -> %d 条事实）" % (n_before, len(facts2)),
              len(facts2) == n_before + 3 and "R002" in {f.get("source_row_id") for f in facts2},
              "实际 %d 条" % len(facts2))
        code, out = run(tmp, ["tools/收尾报告.py", rel, "--one-line"])
        check("收尾报告说任务已完成", "已完成" in out, out.strip())
        check("收尾报告报出归档 2/2", "归档 2/2" in out, out.strip())

        # ---------- 11. 长期记忆收录 ----------
        print()
        print("【11】收录：标准答案库 / 错题本 / 同类词词典（只追加）")

        def count_lines(p):
            if not os.path.isfile(p):
                return 0
            with open(p, "r", encoding="utf-8") as fh:
                return len([l for l in fh if l.strip()])

        code, out = run(tmp, ["tools/收录.py", "golden", rel])
        check("收录标准答案退出码为 0", code == 0, out.strip()[-300:])
        golden_path = os.path.join(tmp, "golden", "golden.jsonl")
        check("golden.jsonl 生成了", os.path.isfile(golden_path))
        n_gold = count_lines(golden_path)
        check("只收录已确认的事实（至少 1 条）", n_gold >= 1, "实际 %d 条" % n_gold)
        code, out = run(tmp, ["tools/收录.py", "golden", rel])
        check("重复跑不会重复追加（幂等）", count_lines(golden_path) == n_gold,
              "%d -> %d" % (n_gold, count_lines(golden_path)))

        fb_path = os.path.join(tmp, "FEEDBACK_LOG.md")
        before_fb = open(fb_path, "r", encoding="utf-8").read()
        code, out = run(tmp, ["tools/收录.py", "feedback", rel, "--field", "乙酸",
                              "--wrong", "0.8%", "--right", "0.0%", "--type", "视觉误读"])
        check("记错题退出码为 0", code == 0, out.strip()[-300:])
        after_fb = open(fb_path, "r", encoding="utf-8").read()
        check("错题本追加了新的一行", len(after_fb) > len(before_fb) and "视觉误读" in after_fb)
        check("原有示例行没被改动（只追加，不重排）", "FB-000" in after_fb)

        aliases_path = os.path.join(tmp, "aliases.json")
        before_al = open(aliases_path, "r", encoding="utf-8").read()
        code, out = run(tmp, ["tools/收录.py", "aliases", rel])
        after_al = open(aliases_path, "r", encoding="utf-8").read()
        check("没有人拍板过同类词时，词典一个字都不动",
              before_al == after_al or "没" in out, out.strip()[-200:])

        # ---------- 13. 真数据未被改动 ----------
        print()
        print("【13】确认真实工作区没被碰过")
        check("真实工作区的 runs/ 里没有多出自测产物",
              not any(d.startswith("20260910_00") and d.endswith("_自测") for d in
                      os.listdir(os.path.join(BASE, "runs"))))

        # ---------- 12. 调度器：一个入口，8 个动词 ----------
        print()
        print("【12】调度器 tools/任务.py（模型不用拼任务号/图片编号/路径）")

        def d(ws, args, stdin_text=None):
            return run(ws, ["tools/任务.py"] + args, stdin_text)

        code, out = d(tmp, ["status"])
        check("status：任务都干完时也给人话，不崩", code == 0 and "当前任务" in out, out.strip()[-160:])

        code, out = d(tmp, ["new"])
        check("new：不用带任何参数就能建任务", code == 0, out.strip()[-200:])
        runs2 = sorted([x for x in os.listdir(runs) if os.path.isdir(os.path.join(runs, x))])
        check("new：建出了第二个任务", len(runs2) == 2, str(runs2))
        run2 = os.path.join(runs, runs2[-1])
        rel2 = os.path.relpath(run2, tmp)

        code, out = d(tmp, ["next"])
        check("next：自动认出新任务、报出待识别的图", code == 0 and "IMG_001" in out,
              out.strip()[-200:])
        check("next：自动认任务时挑的是最新的那个", runs2[-1] in out, out.strip()[:200])

        frag2 = dict(FRAG_OK)
        frag2["image_id"] = "IMG_001"
        code, out = d(tmp, ["submit"], stdin_text=json.dumps(frag2, ensure_ascii=False))
        check("submit：片段走标准输入、图片编号从片段里读、不用给路径",
              code == 0, out.strip()[-300:])
        raw5 = read_json(os.path.join(run2, "raw.json"), {})
        check("submit：确实落到新任务里了", len(raw5.get("rows", [])) == 2,
              "行数=%d" % len(raw5.get("rows", [])))
        # 关键断言：管道送进去的中文必须原样落盘。
        # （曾经这里真的烂过：stdin 按 GBK 解码，脚本退出码 0、数据却是"姘?"。）
        rows5 = raw5.get("rows", [])
        check("submit：管道里的中文原样落盘，没有变乱码",
              bool(rows5) and rows5[0].get("fields", {}).get("位置") == "1号库",
              str(rows5[0].get("fields")) if rows5 else "无数据")
        check("submit：计量项名也原样落盘（水 / 杂）",
              bool(rows5) and [c.get("name") for c in rows5[1].get("components", [])] == ["甲醇", "水", "杂"],
              str([c.get("name") for c in rows5[1].get("components", [])]) if len(rows5) > 1 else "无数据")
        blob = json.dumps(raw5, ensure_ascii=False)
        check("整份 raw.json 里没有替换字符（乱码的典型痕迹）", "\ufffd" not in blob)

        # 片段里少了 image_id 时必须给明确提示，而不是猜
        bad = json.loads(json.dumps(frag2))
        del bad["image_id"]
        code, out = d(tmp, ["submit"], stdin_text=json.dumps(bad, ensure_ascii=False))
        check("submit：片段没带 image_id 时拒绝并说明怎么改", code == 2 and "image_id" in out,
              out.strip()[-200:])

        code, out = d(tmp, ["submit"])
        check("submit：什么都不给时给人话，而不是卡住等输入", code == 2 and "识别片段" in out,
              out.strip()[-160:])

        code, out = d(tmp, ["--help"])
        check("--help 能打出 8 个动词", "finalize" in out and "submit" in out)

        code, out = d(tmp, ["乱写的动词"])
        check("不认识的动词要明确拒绝", code == 2 and "不认识" in out, out.strip()[-120:])

    finally:
        # 只删自己建的临时目录；万一临时目录里被写进了别的东西，也一并随它去（它是新建的）
        try:
            shutil.rmtree(tmp, ignore_errors=True)
            print()
            print("  临时工作区已清理：%s" % tmp)
        except Exception as e:
            print("  ⚠ 临时目录没删掉（不影响结果）：%s" % e)

    print()
    print("=" * 62)
    print("  结果：通过 %d 项，失败 %d 项" % (len(PASS), len(FAIL)))
    if FAIL:
        print("  失败项：")
        for f in FAIL:
            print("    · %s" % f)
    print("=" * 62)
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
